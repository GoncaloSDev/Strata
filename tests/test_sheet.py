"""Folha Excel: modelo (D-05..D-07) e leitura da folha do utilizador (D-08, D-09, LIST-05)."""

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from strata import sheet
from strata.__main__ import main

HEADERS = ("Produto", "Família", "Tamanho", "Unidade", "Qty", "Referência", "Notas")


# --- modelo -----------------------------------------------------------------


def test_modelo_cabecalhos_e_exemplos(tmp_path):
    path = tmp_path / "modelo.xlsx"
    sheet.write_template(path)
    rows = list(load_workbook(path).active.iter_rows(min_row=1, max_row=3, values_only=True))
    assert rows[0] == HEADERS
    assert rows[1][:5] == ("OSB 12 mm", "Madeira", "2500x1250x12 mm", "un", 30)
    assert (rows[2][0], rows[2][3], rows[2][4]) == ("cimento", "saco", 10)


def test_modelo_lista_pendente_da_unidade(tmp_path):
    path = tmp_path / "modelo.xlsx"
    sheet.write_template(path)
    dvs = load_workbook(path).active.data_validations.dataValidation
    assert len(dvs) == 1
    assert dvs[0].type == "list"
    assert dvs[0].formula1 == '"un,m²,m,kg,L,saco"'
    assert str(dvs[0].sqref) == "D2:D1000"


def test_modelo_nunca_sobrescreve(tmp_path):
    path = tmp_path / "modelo.xlsx"
    sheet.write_template(path)
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        sheet.write_template(path)
    assert path.read_bytes() == before


def test_modelo_comando_duas_vezes(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    assert main(["modelo"]) == 0
    assert (tmp_path / "modelo.xlsx").exists()
    before = (tmp_path / "modelo.xlsx").read_bytes()
    capsys.readouterr()
    assert main(["modelo"]) == 1
    assert "modelo.xlsx já existe nesta pasta; não foi alterado." in capsys.readouterr().out
    assert (tmp_path / "modelo.xlsx").read_bytes() == before


def test_modelo_lido_pela_app(tmp_path):
    path = tmp_path / "modelo.xlsx"
    sheet.write_template(path)
    items, invalid = sheet.read_sheet(path)
    assert invalid == []
    assert [i.unit for i in items] == ["un", "saco"]
    assert items[0].family == "Madeira"
    assert items[0].qty == Decimal(30)
