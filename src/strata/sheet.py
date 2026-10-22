"""Ler a folha Excel do utilizador (só leitura) e gravar o modelo vazio."""

import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from strata.models import Invalida, Item
from strata.text import UNITS, norm, to_decimal


class SheetError(Exception):
    pass


REQUIRED = ("produto", "unidade", "qty")
OPTIONAL = ("familia", "tamanho", "referencia", "notas")
HEADERS = ("Produto", "Família", "Tamanho", "Unidade", "Qty", "Referência", "Notas")
UNIT_CHOICES = "un,m²,m,kg,L,saco"


def find_table(wb):
    """Primeira folha com Produto, Unidade e Qty nas linhas 1-10 (D-08, D-09).

    Devolve (folha, linha do cabeçalho, {nome normalizado: índice da coluna}).
    """
    best_missing = list(REQUIRED)
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            cols = {}
            for i, v in enumerate(row):
                if v is not None and str(v).strip():
                    cols.setdefault(norm(v), i)  # coluna repetida: vale a primeira
            missing = [h for h in REQUIRED if h not in cols]
            if not missing:
                return ws, r, cols
            if len(missing) < len(best_missing):
                best_missing = missing
    raise SheetError(
        "nenhuma folha tem os cabeçalhos obrigatórios: faltam " + ", ".join(best_missing)
    )


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _text(v) -> str | None:
    return None if _blank(v) else str(v).strip()


def read_sheet(path) -> tuple[list[Item], list[Invalida]]:
    path = Path(path)
    try:
        wb = load_workbook(path, data_only=True)  # modo normal: números de linha certos (Pitfall 7)
    except FileNotFoundError:
        raise SheetError(f"ficheiro não encontrado: {path}") from None
    except PermissionError:
        raise SheetError(f"sem permissão para ler {path} (está aberto noutro programa?)") from None
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise SheetError(f"não consegui abrir {path.name}: é .xls ou não é um .xlsx válido (guarde a folha como .xlsx)") from None

    ws, header, cols = find_table(wb)

    def cell(row, name):
        i = cols.get(name)
        return row[i] if i is not None and i < len(row) else None

    items, invalid = [], []
    for n, row in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), start=header + 1):
        if all(_blank(v) for v in row):
            continue
        reasons = []
        product = _text(cell(row, "produto"))
        if product is None:
            reasons.append("Produto vazio")

        raw_unit = cell(row, "unidade")
        unit = None
        if _blank(raw_unit):
            reasons.append("Unidade vazia")
        else:
            unit = UNITS.get(norm(raw_unit))
            if unit is None:
                reasons.append(f"unidade desconhecida: {str(raw_unit).strip()}")

        raw_qty = cell(row, "qty")
        qty = to_decimal(raw_qty)
        if _blank(raw_qty):
            reasons.append("Qty vazia")
        elif qty is None or qty <= 0:
            reasons.append(f"Qty não é um número positivo: {str(raw_qty).strip()}")

        if reasons:
            invalid.append(Invalida(n, "; ".join(reasons)))
            continue
        items.append(Item(
            row=n, product=product, qty=qty, unit=unit,
            family=_text(cell(row, "familia")),
            size=_text(cell(row, "tamanho")),
            reference=_text(cell(row, "referencia")),
            notes=_text(cell(row, "notas")),
        ))
    return items, invalid


def write_template(path) -> None:
    """Grava o modelo (D-05, D-06, D-07). Nunca sobrescreve: FileExistsError sobe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(list(HEADERS))
    ws.append(["OSB 12 mm", "Madeira", "2500x1250x12 mm", "un", 30, None, None])
    ws.append(["cimento", None, None, "saco", 10, None, None])
    dv = DataValidation(
        type="list", formula1=f'"{UNIT_CHOICES}"', allow_blank=True, showErrorMessage=True,
        errorTitle="Unidade", error="Escolha da lista: " + UNIT_CHOICES.replace(",", ", "),
    )
    ws.add_data_validation(dv)
    dv.add("D2:D1000")
    for col, width in zip("ABCDEFG", (30, 16, 18, 10, 8, 16, 30)):
        ws.column_dimensions[col].width = width
    with open(path, "xb") as f:  # criação exclusiva: um modelo existente nunca é tocado
        wb.save(f)
